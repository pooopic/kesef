"""
ממיר קובץ אקסל של פירוט כרטיס אשראי לפורמט הייבוא של האפליקציה (docs/IMPORT_FORMAT.md).

שימוש:
    python tools/convert_card_excel.py "סיכום הוצאות 2026.xlsx"

הפלט: קובץ "ייבוא - <שם הקובץ>.csv" באותה תיקייה. אותו מעבירים לטלפון ומייבאים מההגדרות.

הקובץ צריך לכלול גיליון עם העמודות:
חודש, כרטיס, בית עסק, קטגוריה, תאריך עסקה, סכום העסקה, סוג העסקה, פירוט, תאריך החיוב, סכום החיוב
"""

import csv
import datetime
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

# הקטגוריות של חברת האשראי → הקטגוריות של האפליקציה.
# קטגוריה שלא מופיעה כאן נשארת בשם המקורי שלה (ותיווצר באפליקציה בייבוא).
CATEGORY_MAP = {
    "סופרמרקט ומזון": "מזון וסופר",
    "מסעדות, בתי קפה ומאפיות": "אוכל בחוץ",
    "דלק, חניה ותחבורה": "תחבורה ורכב",
    "פנאי, בילוי ותיירות": "פנאי וחופשות",
    "קניות כלליות": "קניות ובית",
    "בריאות ותרופות": "בריאות וטיפוח",
    "חשמל ותשתיות": "דיור",
    "תקשורת ודיגיטל": "מנויים וחשבונות",
    "העברות כספים": "העברות ומזומן",
    "רכב ותחזוקה": "תחבורה ורכב",
    "אופנה, הלבשה והנעלה": "קניות ובית",
    "ספורט": "פנאי וחופשות",
    "איבזור הבית": "קניות ובית",
    "תרומות": "מתנות ותרומות",
    "ביטוח ומיסים": "ביטוח, מיסים ועמלות",
    "חשמל ואלקטרוניקה": "קניות ובית",
    "אחר": "אחר",
}

COLUMNS = ["חודש", "כרטיס", "בית עסק", "קטגוריה", "תאריך עסקה", "סכום העסקה", "סוג העסקה", "פירוט", "תאריך החיוב", "סכום החיוב"]
OUT_HEADER = ["תאריך", "סכום", "סוג", "קטגוריה", "תיאור", "אמצעי תשלום", "הערה"]


def map_category(name):
    name = (name or "").strip()
    if name.startswith("מנוי"):  # "מנוי - Spotify", "מנוי - פרטנר"...
        return "מנויים וחשבונות"
    return CATEGORY_MAP.get(name, name or "אחר")


def to_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, str):
        try:
            return datetime.datetime.strptime(value.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def find_sheet(wb):
    for ws in wb.worksheets:
        header = [str(c.value).strip() if c.value else "" for c in ws[1]]
        if all(col in header for col in ("בית עסק", "סכום החיוב", "תאריך עסקה")):
            return ws, header
    sys.exit("לא נמצא גיליון עם העמודות הנדרשות: " + ", ".join(COLUMNS))


def convert(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, header = find_sheet(wb)
    col = {name: header.index(name) for name in COLUMNS if name in header}

    out_rows, skipped, categories = [], [], Counter()
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda name: row[col[name]] if name in col else None
        charge = get("סכום החיוב")
        if not isinstance(charge, (int, float)) or charge == 0:
            skipped.append(i)  # שורת סיכום, הערה, או עסקה שעוד לא חויבה
            continue

        kind = (get("סוג העסקה") or "").strip()
        detail = (get("פירוט") or "").strip()
        is_installment = kind.startswith("תשלום") or kind.startswith("קרדיט")

        # עסקה רגילה: לפי תאריך העסקה. תשלומים: לפי תאריך החיוב, כי כל חודש יורד רק חלק.
        date = to_date(get("תאריך החיוב")) if is_installment else to_date(get("תאריך עסקה"))
        date = date or to_date(get("תאריך עסקה")) or to_date(get("תאריך החיוב"))
        if date is None:
            skipped.append(i)
            continue

        description = (get("בית עסק") or "").strip()
        notes = []
        m = re.search(r"(\d+)\s*מ\s*-\s*(\d+)", detail)
        if is_installment and m:
            notes.append(f"תשלום {m.group(1)} מתוך {m.group(2)} (סה\"כ {get('סכום העסקה')})")
        transfer = re.search(r"העברה ל\s*(.+?)(?:\s+עבור\s+(.+))?$", detail)
        if description.upper() in ("BIT", "PAYBOX") and transfer:
            description = f"{description} · {transfer.group(1).strip()}"
            if transfer.group(2):
                notes.append(transfer.group(2).strip())
        elif kind == 'חו"ל':
            notes.append('חו"ל')

        category = map_category(get("קטגוריה"))
        categories[category] += 1
        out_rows.append([
            date.strftime("%d/%m/%Y"),
            f"{charge:.2f}".rstrip("0").rstrip("."),
            "הוצאה",
            category,
            description,
            (get("כרטיס") or "").strip(),
            " · ".join(notes),
        ])

    out_path = Path(path).with_name(f"ייבוא - {Path(path).stem}.csv")
    # utf-8-sig = UTF-8 עם BOM, כדי שגם Excel יפתח את הקובץ בעברית תקינה
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(OUT_HEADER)
        writer.writerows(out_rows)

    total = sum(float(r[1]) for r in out_rows)
    print(f"נכתב: {out_path}")
    print(f"{len(out_rows)} תנועות, סה\"כ {total:,.2f} ש\"ח. דולגו {len(skipped)} שורות (סיכום/0).")
    print("קטגוריות:")
    for name, count in categories.most_common():
        print(f"  {name}: {count}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    convert(sys.argv[1])
